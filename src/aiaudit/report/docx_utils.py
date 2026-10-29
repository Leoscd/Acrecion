"""Helpers de python-docx. Si existe report/templates/base.docx se usa como base (logo, estilos, encabezado de la consultora)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

BASE = Path(__file__).parent / "templates" / "base.docx"
COLOR = RGBColor(0x1F, 0x3A, 0x5F)


def nuevo(titulo: str, subtitulo: str):
    doc = Document(str(BASE)) if BASE.exists() else Document()
    estilo = doc.styles["Normal"]
    estilo.font.name, estilo.font.size = "Calibri", Pt(10.5)
    t = doc.add_heading(titulo, 0)
    for run in t.runs:
        run.font.color.rgb = COLOR
    p = doc.add_paragraph(f"{subtitulo}\n{date.today():%d/%m/%Y}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return doc


def tabla(doc, encabezados: list[str], filas: list[list], anchos: list[float] | None = None):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Light Grid Accent 1"
    for celda, texto in zip(t.rows[0].cells, encabezados):
        celda.text = str(texto)
        celda.paragraphs[0].runs[0].bold = True
    for fila in filas:
        for celda, valor in zip(t.add_row().cells, fila):
            celda.text = "" if valor is None else str(valor)
    doc.add_paragraph()
    return t


def vinetas(doc, items: list[str], vacio: str = "Sin datos."):
    for i in items or [vacio]:
        doc.add_paragraph(str(i), style="List Bullet")


def dinero(valor: float, moneda: str) -> str:
    return f"{moneda} {valor:,.0f}".replace(",", ".")
